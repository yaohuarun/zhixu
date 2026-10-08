"""Generate small reproducible parsing/OCR documents; no real customer material."""
from pathlib import Path
import io

import fitz
from docx import Document
from PIL import Image, ImageDraw, ImageFont


def main():
    target = Path("data/acceptance")
    target.mkdir(parents=True, exist_ok=True)
    content = "知识库支持 PDF、Word 和 TXT。扫描件通过本地 OCR 处理。"
    (target / "guide.txt").write_text(content, encoding="utf-8")
    word = Document()
    word.add_heading("知识库操作说明", level=1)
    word.add_paragraph(content)
    word.add_paragraph("接入方式如下：")
    word.add_paragraph("上传文件或文件夹", style="List Bullet")
    word.add_paragraph("指定服务器目录", style="List Bullet")
    table = word.add_table(rows=2, cols=2)
    for cell, text in zip([c for row in table.rows for c in row.cells], ["格式", "支持", "PDF", "扫描件 OCR"]):
        cell.text = text
    word.save(target / "guide.docx")
    fonts = [Path("C:/Windows/Fonts/msyh.ttc"),
             Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")]
    font_path = next((p for p in fonts if p.exists()), None)
    if not font_path:
        raise RuntimeError("Install a CJK font for the scanned PDF fixture")
    image = Image.new("RGB", (1200, 900), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(font_path), 42)
    for index, line in enumerate(["本地扫描件识别", "知识库支持扫描件 PDF。", "文字通过本地 OCR 提取。",
                                  "Local OCR recognizes scanned documents."]):
        draw.text((60, 80 + 100 * index), line, font=font, fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    mixed = fitz.open()
    page = mixed.new_page()
    page.insert_text((70, 90), "Knowledge base supports PDF, Word and TXT.", fontsize=14)
    page = mixed.new_page(width=600, height=450)
    page.insert_image(page.rect, stream=buffer.getvalue())
    mixed.save(target / "mixed.pdf")
    scanned = fitz.open()
    scanned.insert_pdf(mixed, from_page=1, to_page=1)
    scanned.save(target / "scanned.pdf")
    text = fitz.open()
    text.insert_pdf(mixed, from_page=0, to_page=0)
    text.save(target / "text.pdf")
    print(target.resolve())


if __name__ == "__main__":
    main()

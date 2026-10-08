import asyncio
import json

import fitz
import httpx
import pytest
from docx import Document
from pydantic import SecretStr

from app.chunking import Counter, chunk_elements
from app.config import Settings, settings
from app.errors import AppError, safe_error
from app.ingestion import enumerate_directory, relative_path, source_directory, stable_snapshot
from app.parsing import parse_direct
from app.providers import Providers
from app.retrieval import rrf, sanitize_citations


def test_path_guards(tmp_path, monkeypatch):
    monkeypatch.setattr(settings(), "sources_root", tmp_path)
    assert relative_path("folder\\doc.txt") == "folder/doc.txt"
    for path in ("../secret.txt", "/secret.txt", "C:\\secret.txt", "a/../b", "a\x00.txt"):
        with pytest.raises(AppError):
            relative_path(path)
    with pytest.raises(AppError):
        source_directory(str(tmp_path.parent / "elsewhere"))
    (tmp_path / "nested").mkdir()
    (tmp_path / "nested/a.txt").write_text("a")
    assert list(enumerate_directory(tmp_path)) == ["nested/a.txt"]


def test_stable_snapshot_and_size(tmp_path, monkeypatch):
    monkeypatch.setattr(settings(), "storage_root", tmp_path / "store")
    file = tmp_path / "a.txt"
    file.write_text("知识库", encoding="utf-8")
    snapshot, digest = stable_snapshot(file, ".txt")
    assert snapshot.read_bytes() == file.read_bytes() and len(digest) == 64
    monkeypatch.setattr(settings(), "max_upload_bytes", 2)
    with pytest.raises(AppError, match="大小限制"):
        stable_snapshot(file, ".txt")


def test_parsers_and_provenance(tmp_path):
    txt = tmp_path / "manual.txt"
    txt.write_bytes("中文段落。\n\n支持 TXT。".encode("gb18030"))
    assert "中文" in parse_direct(txt)[0]["text"]
    doc = Document()
    doc.add_heading("产品说明", 1)
    doc.add_paragraph("支持 PDF、Word 和 TXT。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "格式", "支持"
    table.cell(1, 0).text, table.cell(1, 1).text = "PDF", "是"
    word = tmp_path / "manual.docx"
    doc.save(word)
    elements = parse_direct(word)
    assert elements[-1]["kind"] == "table" and elements[-1]["section"] == ["产品说明"]
    assert all(e["page"] is None for e in elements)
    pdf = fitz.open()
    for text in ("Knowledge manual\nSupported documents include PDF and Word.", "Incremental updates are supported."):
        page = pdf.new_page()
        page.insert_text((70, 70), text)
    pdf.new_page()  # A truly blank page must not require initialized OCR assets.
    path = tmp_path / "manual.pdf"
    pdf.save(path)
    pdf.close()
    parsed = parse_direct(path)
    assert {e["page"] for e in parsed} == {1, 2}
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"not a PDF")
    with pytest.raises(AppError, match="损坏"):
        parse_direct(broken)


def test_chunks_sections_overlap_and_tables(tmp_path):
    txt = tmp_path / "long.txt"
    txt.write_text("第一节内容。" * 400, encoding="utf-8")
    chunks = chunk_elements(parse_direct(txt), "manual.txt")
    counter = Counter()
    assert len(chunks) > 5 and all(counter.count(c.retrieval_text) <= 900 for c in chunks)
    assert chunks[0].provenance["end"] > chunks[1].provenance["start"]
    elements = [
        {"text": "A", "kind": "paragraph", "section": ["A"], "page": 1, "start": 0, "end": 1, "element_index": 0},
        {"text": "B", "kind": "paragraph", "section": ["B"], "page": 2, "start": 2, "end": 3, "element_index": 1},
    ]
    assert len(chunk_elements(elements, "test")) == 2
    rows = ["字段 | 描述"] + [f"{i} | " + "描述" * 20 for i in range(20)]
    table = {**elements[0], "kind": "table", "text": "\n".join(rows), "rows": rows}
    drafts = chunk_elements([table], "table")
    assert len(drafts) > 1 and all(c.original.startswith(rows[0]) for c in drafts)
    assert all(c.provenance["counter"] == "utf8-byte-upper-budget-v1" for c in drafts)


def configured():
    return Settings(_env_file=None, embedding_url="https://model.test/embed", embedding_key=SecretStr("private-key"),
                    rerank_url="https://model.test/rerank", rerank_key=SecretStr("private-key"),
                    deepseek_url="https://model.test/chat", deepseek_model="test-model",
                    deepseek_key=SecretStr("private-key"), model_retries=0)


@pytest.mark.parametrize("index_field", ["text_index", "index"])
def test_embedding_contract_orders_and_rejects_dimension(index_field):
    config = configured()
    def reply(request):
        body = json.loads(request.content)
        assert body["model"] == "qwen3.7-text-embedding-flash"
        assert body["parameters"]["text_type"] == "query"
        return httpx.Response(200, json={"output": {"embeddings": [
            {index_field: 1, "embedding": [0.2] * 1024}, {index_field: 0, "embedding": [0.1] * 1024},
        ]}})
    provider = Providers(config, httpx.MockTransport(reply))
    vectors = asyncio.run(provider.embed(["A", "B"], config.embedding_config(), query=True))
    assert vectors[0][0] == 0.1 and vectors[1][0] == 0.2
    provider = Providers(config, httpx.MockTransport(lambda r: httpx.Response(200, json={"output": {
        "embeddings": [{"text_index": 0, "embedding": [1]}]}})))
    with pytest.raises(AppError, match="维度"):
        asyncio.run(provider.embed(["A"], config.embedding_config()))


def test_rerank_invalid_indexes_and_error_redaction():
    provider = Providers(configured(), httpx.MockTransport(lambda r: httpx.Response(200, json={
        "output": {"results": [{"index": 4, "relevance_score": 0.8}]}})))
    with pytest.raises(AppError, match="索引"):
        asyncio.run(provider.rerank("问题", ["资料"]))
    assert "private-key" not in safe_error(RuntimeError("Bearer private-key"))
    missing = Providers(Settings(_env_file=None))
    with pytest.raises(AppError, match="未配置"):
        missing.credentials("deepseek")


def test_rrf_and_citations():
    fused = rrf([[('a', 0.8), ('b', 0.7)], [('b', 0.1), ('c', 0.9)]], 60)
    assert fused[0][0] == 'b' and len(fused) == 3
    assert sanitize_citations("依据 [1]，虚构 [99]", {1}) == "依据 [1]，虚构 "


def test_stream_requires_terminal_marker():
    cfg = configured()
    provider = Providers(cfg, httpx.MockTransport(lambda r: httpx.Response(200, text=
        'data: {"choices":[{"delta":{"content":"hello"}}]}\n\n', headers={"content-type": "text/event-stream"})))
    async def run():
        return [x async for x in provider.stream([{"role": "user", "content": "hi"}])]
    with pytest.raises(AppError, match="异常结束"):
        asyncio.run(run())


def test_rewrite_disables_thinking_and_reports_truncation():
    def reply(request):
        body = json.loads(request.content)
        assert body["thinking"] == {"type": "disabled"}
        return httpx.Response(200, json={"choices": [{"finish_reason": "length",
            "message": {"content": "", "reasoning_content": "unnecessary reasoning"}}]})
    provider = Providers(configured(), httpx.MockTransport(reply))
    with pytest.raises(AppError) as error:
        asyncio.run(provider.chat([{"role": "user", "content": "rewrite"}]))
    assert error.value.code == "output_truncated"


@pytest.mark.parametrize("thinking", ["disabled", "enabled"])
def test_answer_thinking_is_explicit_and_configurable(thinking):
    config = configured()
    config.deepseek_thinking = thinking
    def reply(request):
        body = json.loads(request.content)
        assert body["thinking"] == {"type": thinking}
        assert body["max_tokens"] == 4096
        return httpx.Response(200, text='data: {"choices":[{"delta":{"content":"answer"}}]}\n\ndata: [DONE]\n\n',
                              headers={"content-type": "text/event-stream"})
    provider = Providers(config, httpx.MockTransport(reply))
    async def collect():
        return [piece async for piece in provider.stream([{"role": "user", "content": "question"}])]
    assert asyncio.run(collect()) == ["answer"]


def test_list_introduction_stays_with_items(monkeypatch):
    monkeypatch.setattr(settings(), "chunk_target", 200)
    elements = [
        {"text": "a" * 210 + ":", "kind": "paragraph", "section": [], "page": None,
         "start": 0, "end": 211, "element_index": 0},
        {"text": "- first item", "kind": "list", "section": [], "page": None,
         "start": 213, "end": 225, "element_index": 1},
    ]
    drafts = chunk_elements(elements, "list")
    assert len(drafts) == 1 and "first item" in drafts[0].original


def test_readiness_does_not_hide_missing_dependency(monkeypatch):
    from fastapi.testclient import TestClient
    from app.main import app
    monkeypatch.setattr("app.main._database_ready", lambda: None)
    def missing():
        raise RuntimeError("connection refused")
    monkeypatch.setattr("app.vectors.vector_store", missing)
    response = TestClient(app).get("/api/health/ready")
    assert response.status_code == 503
    assert response.json()["dependencies"] == {"postgres": "ok", "qdrant": "unavailable"}


def test_scanned_page_with_native_header_still_uses_ocr(tmp_path, monkeypatch):
    from PIL import Image
    import io
    buffer = io.BytesIO()
    Image.new("RGB", (600, 800), "white").save(buffer, format="PNG")
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_image(page.rect, stream=buffer.getvalue())
    page.insert_text((70, 60), "Native header does not represent the scanned body")
    path = tmp_path / "scan-with-header.pdf"
    pdf.save(path)
    pdf.close()
    class Result:
        json = {"res": {"parsing_res_list": [{"block_content": "中文断行\n合并文本", "block_label": "text",
                                             "block_bbox": [0, 0, 300, 100]}]}}
    class OCR:
        def predict(self, path):
            return [Result()]
    monkeypatch.setattr("app.parsing.ocr_engine", OCR)
    parsed = parse_direct(path)
    assert parsed[0]["extraction"] == "ocr" and parsed[0]["text"] == "中文断行合并文本"
    assert parsed[0]["raw_text"] == "中文断行\n合并文本" and parsed[0]["page"] == 1
